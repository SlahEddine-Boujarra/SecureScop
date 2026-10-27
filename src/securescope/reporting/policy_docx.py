"""Editable Word export of the current structured policy document."""

from io import BytesIO

from docx import Document

from securescope.policies.documents import validate_blocks


def policy_docx(row: dict, blocks: list[dict]) -> bytes:
    """Export real headings, paragraphs and emphasis to a DOCX document.

    Args:
        row: Saved policy metadata.
        blocks: Structured policy content.
    Returns:
        DOCX bytes suitable for a download.
    Raises:
        ValueError: If document content is invalid.
    """
    document = Document()
    document.add_heading(str(row.get("title", "Policy")), 0)
    document.add_paragraph(
        f"Version {row.get('policy_version', 0)} · {row.get('status', 'Draft')}"
        f"\nOwner: {row.get('owner', '')} · Review: {row.get('review_date', '')}"
    )
    for block in validate_blocks(blocks):
        kind = block["kind"]
        paragraph = (
            document.add_heading("", level=int(kind[1]))
            if kind.startswith("h")
            else document.add_paragraph(
                style={"bullet": "List Bullet", "number": "List Number"}.get(kind)
            )
        )
        for span in block["runs"]:
            run = paragraph.add_run(span["text"])
            run.bold = bool(span.get("bold"))
            run.italic = bool(span.get("italic"))
    output = BytesIO()
    document.save(output)
    return output.getvalue()

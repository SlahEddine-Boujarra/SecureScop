"""Printable ISO control matrix and evidence package generation."""

from __future__ import annotations

import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime, timezone
from typing import Any, Mapping, Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table as ExcelTable
from openpyxl.worksheet.table import TableStyleInfo
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from securescope.evidence.storage import read_artifact
from securescope.grc.control_catalog import CATALOG_NOTICE, FRAMEWORK_REFERENCE
from securescope.persistence.db_service import get_client_id
from securescope.persistence.evidence_repository import (
    list_evidence_audit_events,
    list_evidence_records,
)
from securescope.persistence.grc_repository import list_control_overview

SAFE_ARCHIVE_NAME = re.compile(r"[^A-Za-z0-9._-]+")
CONFIDENTIALITY = "Confidential - supplied for authorized assessment and audit use"


class AuditorPackageError(RuntimeError):
    """Raised when an auditor deliverable cannot be produced."""


def _text(value: Any) -> str:
    """Return a predictable printable representation."""

    if value is None:
        return ""
    if isinstance(value, (list, dict)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def _excel_value(value: Any) -> Any:
    """Preserve sortable Excel date values and normalize timezone-aware timestamps."""

    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).replace(tzinfo=None)
    if isinstance(value, date):
        return value
    return _text(value)


def _footer(canvas: Any, document: Any) -> None:
    """Draw confidentiality and page numbers on each PDF page."""

    canvas.saveState()
    canvas.setFont("Helvetica", 7)
    canvas.setFillColor(colors.HexColor("#56616f"))
    canvas.drawString(15 * mm, 8 * mm, CONFIDENTIALITY)
    canvas.drawRightString(
        landscape(A4)[0] - 15 * mm,
        8 * mm,
        f"Page {document.page}",
    )
    canvas.restoreState()


def build_control_matrix_pdf(
    client_name: str,
    *,
    prepared_by: str,
    assessment_scope: str,
    rows: Sequence[Mapping[str, Any]] | None = None,
) -> bytes:
    """Build a landscape, multi-page ISO control matrix for printing."""

    controls = list(rows) if rows is not None else list_control_overview(client_name)
    output = io.BytesIO()
    styles = getSampleStyleSheet()
    small = ParagraphStyle("small", parent=styles["BodyText"], fontSize=6.5, leading=8)
    heading = ParagraphStyle(
        "heading",
        parent=styles["Title"],
        textColor=colors.HexColor("#17365d"),
        alignment=TA_CENTER,
    )
    document = SimpleDocTemplate(
        output,
        pagesize=landscape(A4),
        leftMargin=12 * mm,
        rightMargin=12 * mm,
        topMargin=13 * mm,
        bottomMargin=14 * mm,
        title=f"{client_name} ISO control assessment matrix",
        author=prepared_by,
    )
    story: list[Any] = [
        Paragraph("ISO/IEC 27001 Assessment - Control Matrix", heading),
        Spacer(1, 3 * mm),
        Paragraph(f"<b>Workspace:</b> {_text(client_name)}", styles["BodyText"]),
        Paragraph(f"<b>Assessment scope:</b> {_text(assessment_scope)}", styles["BodyText"]),
        Paragraph(f"<b>Prepared by:</b> {_text(prepared_by)}", styles["BodyText"]),
        Paragraph(
            f"<b>Generated:</b> {datetime.now(timezone.utc).isoformat()}", styles["BodyText"]
        ),
        Paragraph(f"<b>Framework:</b> {FRAMEWORK_REFERENCE}", styles["BodyText"]),
        Paragraph(CATALOG_NOTICE, small),
        Spacer(1, 4 * mm),
    ]
    header = [
        "Control",
        "Short label",
        "Applicable",
        "Method",
        "Status",
        "Evidence",
        "Freshness",
        "Owner",
        "Validation",
        "Last review",
    ]
    table_rows: list[list[Any]] = [[Paragraph(f"<b>{item}</b>", small) for item in header]]
    for row in controls:
        table_rows.append(
            [
                Paragraph(_text(row.get("control_id")), small),
                Paragraph(_text(row.get("control_name")), small),
                Paragraph(_text(row.get("applicability") or "Not reviewed"), small),
                Paragraph(
                    _text(row.get("measurement_model") or "checklist").replace("_", " ").title(),
                    small,
                ),
                Paragraph(_text(row.get("assessment_status") or "Not assessed"), small),
                Paragraph(_text(row.get("approved_evidence_count") or "—"), small),
                Paragraph(_text(row.get("evidence_freshness") or "Unknown"), small),
                Paragraph(_text(row.get("control_owner") or "Unassigned"), small),
                Paragraph(_text(row.get("validation_status") or "Pending"), small),
                Paragraph(_text(row.get("assessed_at") or "—"), small),
            ]
        )
    matrix = Table(
        table_rows,
        repeatRows=1,
        colWidths=[
            18 * mm,
            48 * mm,
            23 * mm,
            27 * mm,
            31 * mm,
            15 * mm,
            24 * mm,
            30 * mm,
            23 * mm,
            28 * mm,
        ],
    )
    matrix.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#17365d")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#b7c2ce")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f5f8fb")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 3),
                ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]
        )
    )
    story.append(matrix)
    document.build(story, onFirstPage=_footer, onLaterPages=_footer)
    return output.getvalue()


def build_control_matrix_xlsx(
    controls: Sequence[Mapping[str, Any]],
    evidence: Sequence[Mapping[str, Any]],
    audit_events: Sequence[Mapping[str, Any]],
) -> bytes:
    """Build a filterable Excel workbook for consultant and auditor analysis."""

    workbook = Workbook()
    control_sheet = workbook.active
    control_sheet.title = "Control Matrix"
    control_headers = [
        "Control",
        "Control name",
        "Applicable",
        "Assessment method",
        "Automation",
        "Status",
        "Evidence freshness",
        "Owner",
        "Validation",
        "Last review",
    ]
    control_sheet.append(control_headers)
    for row in controls:
        control_sheet.append(
            [
                _text(row.get("control_id")),
                _text(row.get("control_name")),
                _text(row.get("applicability")),
                _text(row.get("measurement_model")),
                _text(row.get("automation_level")),
                _text(row.get("assessment_status")),
                _text(row.get("evidence_freshness")),
                _text(row.get("control_owner")),
                _text(row.get("validation_status")),
                _excel_value(row.get("assessed_at")),
            ]
        )
    evidence_sheet = workbook.create_sheet("Validated Evidence")
    evidence_headers = [
        "Evidence ID",
        "Title",
        "Source",
        "Method",
        "Collected",
        "Expires",
        "Owner",
        "Validator",
        "SHA-256",
        "Approved controls",
        "Summary",
    ]
    evidence_sheet.append(evidence_headers)
    for row in evidence:
        approved = [
            item.get("control_id") for item in row.get("controls") or [] if item.get("approved")
        ]
        evidence_sheet.append(
            [
                row.get("id"),
                _text(row.get("title")),
                _text(row.get("tool_id")),
                _text(row.get("collection_method")),
                _excel_value(row.get("collected_at")),
                _excel_value(row.get("expires_at")),
                _text(row.get("owner")),
                _text(row.get("validator")),
                _text(row.get("sha256")),
                ", ".join(approved),
                _text(row.get("summary")),
            ]
        )
    audit_sheet = workbook.create_sheet("Audit History")
    audit_headers = ["Event", "Entity", "Entity ID", "Actor", "Summary", "Timestamp", "Event hash"]
    audit_sheet.append(audit_headers)
    for row in audit_events:
        audit_sheet.append(
            [
                _text(row.get("event_type")),
                _text(row.get("entity_type")),
                _text(row.get("entity_id")),
                _text(row.get("actor")),
                _text(row.get("summary")),
                _excel_value(row.get("created_at")),
                _text(row.get("event_hash")),
            ]
        )
    table_names = {
        "Control Matrix": "ControlMatrixTable",
        "Validated Evidence": "ValidatedEvidenceTable",
        "Audit History": "AuditHistoryTable",
    }
    date_columns = {
        "Control Matrix": ("J",),
        "Validated Evidence": ("E", "F"),
        "Audit History": ("F",),
    }
    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17365D")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(max(len(_text(cell.value)) for cell in column) + 2, 50)
            sheet.column_dimensions[letter].width = max(12, width)
            for cell in column[1:]:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        for letter in date_columns[sheet.title]:
            for cell in sheet[letter][1:]:
                if isinstance(cell.value, (date, datetime)):
                    cell.number_format = "yyyy-mm-dd hh:mm"
        table = ExcelTable(displayName=table_names[sheet.title], ref=sheet.dimensions)
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


def build_auditor_package(
    client_name: str,
    *,
    prepared_by: str,
    assessment_scope: str,
    include_raw_evidence: bool = True,
) -> bytes:
    """Build a ZIP with printable matrix, workbook, register, hashes and raw evidence."""

    if not prepared_by.strip() or not assessment_scope.strip():
        raise AuditorPackageError("Prepared-by and assessment scope are required.")
    controls = list_control_overview(client_name)
    evidence = list_evidence_records(client_name, validation_status="Validated", limit=5_000)
    audit_events = list_evidence_audit_events(client_name, limit=10_000)
    pdf = build_control_matrix_pdf(
        client_name,
        prepared_by=prepared_by,
        assessment_scope=assessment_scope,
        rows=controls,
    )
    xlsx = build_control_matrix_xlsx(controls, evidence, audit_events)
    register = json.dumps(
        {
            "workspace": client_name,
            "assessment_scope": assessment_scope,
            "prepared_by": prepared_by,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "framework": FRAMEWORK_REFERENCE,
            "catalog_notice": CATALOG_NOTICE,
            "validated_evidence": evidence,
        },
        ensure_ascii=False,
        indent=2,
        default=str,
    ).encode("utf-8")
    files: dict[str, bytes] = {
        "reports/ISO-control-matrix.pdf": pdf,
        "reports/ISO-control-matrix.xlsx": xlsx,
        "registers/validated-evidence.json": register,
    }
    if include_raw_evidence:
        client_id = get_client_id(client_name, create=False)
        for row in evidence:
            reference = row.get("storage_reference")
            if not reference:
                continue
            safe = SAFE_ARCHIVE_NAME.sub("_", _text(row.get("original_filename"))).strip("._")
            filename = safe or f"evidence-{row.get('id')}.bin"
            files[f"evidence/{row.get('id')}_{filename}"] = read_artifact(
                client_id, str(reference), expected_sha256=row.get("sha256")
            )
    manifest_lines = [
        f"{hashlib.sha256(content).hexdigest()}  {name}" for name, content in sorted(files.items())
    ]
    files["SHA256SUMS.txt"] = ("\n".join(manifest_lines) + "\n").encode("utf-8")
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name, content in sorted(files.items()):
            archive.writestr(name, content)
    return output.getvalue()

"""Regression tests for the audit, risk and workbook extensions."""

from io import BytesIO
from pathlib import Path

from openpyxl import Workbook, load_workbook

from securescope.grc.document_registers import REGISTER_SCHEMAS, normalize_register_rows
from securescope.grc.risk_register import risk_level, risk_score, starter_risks
from securescope.persistence.audit_scope_repository import default_schedule_rows
from securescope.persistence.migration_service import discover_migrations
from securescope.reporting.iso_assurance_export import build_iso_assurance_workbook
from securescope.reporting.register_exchange import import_register_workbook

ROOT = Path(__file__).resolve().parents[1]


def test_disabled_migrations_are_excluded_from_upgrade_discovery(tmp_path: Path) -> None:
    """An unfinished SQL file must never be promoted by discovery implicitly."""
    (tmp_path / "001_released.sql").write_text("SELECT 1;", encoding="utf-8")
    (tmp_path / "002_unreleased.sql.disabled").write_text("invalid SQL", encoding="utf-8")
    assert [item.version for item in discover_migrations(tmp_path)] == [1]


def test_risk_score_keeps_unassessed_values_blank() -> None:
    """Do not convert incomplete residual-risk data into a low risk."""

    assert risk_score(3, 4) == 12
    assert risk_level(12) == "High"
    assert risk_score(None, 4) is None
    assert risk_level(None) == "Not assessed"


def test_starter_risk_register_matches_the_requested_workflow() -> None:
    """The starter register provides the supplied risk-register shape."""

    rows = starter_risks()
    assert len(rows) == 13
    assert {"pre_likelihood", "pre_impact", "treatment_action", "post_likelihood"} <= rows[0].keys()


def test_iso_workbook_contains_tables_and_native_charts() -> None:
    """Excel export contains dashboards, schedule, SoA and risk charts."""

    workbook_bytes = build_iso_assurance_workbook(
        client_name="Demo",
        gap_rows=[],
        soa_rows=[
            {
                "control_id": "A.5.1",
                "control_name": "Policies",
                "applicability": "Applicable",
                "implementation_status": "Implemented",
            }
        ],
        schedule_rows=default_schedule_rows(),
        risks=starter_risks(),
    )
    workbook = load_workbook(BytesIO(workbook_bytes), data_only=False)
    assert {"Gap Dashboard", "SoA Dashboard", "Audit Schedule", "Risk Register"} <= set(
        workbook.sheetnames
    )
    assert len(workbook["Gap Dashboard"]._charts) == 2
    assert len(workbook["SoA Dashboard"]._charts) == 2
    assert len(workbook["Risk Charts"]._charts) == 1
    assert workbook["Risk Register"]["N5"].value.startswith("=IF(")


def test_unsupported_streamlit_button_column_is_not_used() -> None:
    """Keep the Nuclei/OpenVAS report table compatible with Streamlit 1.62."""

    sources = [
        ROOT / "src" / "securescope" / "app.py",
        ROOT / "src" / "securescope" / "ui" / "dual_engine_ui.py",
    ]
    assert all("ButtonColumn" not in path.read_text(encoding="utf-8") for path in sources)


def test_register_import_accepts_blank_templates_and_scenario_rows() -> None:
    """Read the workbook header below title rows without inventing blank records."""

    blank = Workbook()
    blank_sheet = blank.active
    blank_sheet.title = "Risk Assessment"
    blank_sheet.cell(8, 2, "Scenario-Based Risk Assessment and Treatment Tool")
    fields = list(REGISTER_SCHEMAS["scenario_risks"][1].values())
    for column, label in enumerate(fields, 1):
        blank_sheet.cell(13, column, label)
    blank_output = BytesIO()
    blank.save(blank_output)
    assert import_register_workbook(blank_output.getvalue(), "scenario_risks") == []

    populated = Workbook()
    sheet = populated.active
    sheet.title = "Risk Assessment"
    for column, label in enumerate(fields, 1):
        sheet.cell(13, column, label)
    sheet.cell(14, 1, 1)
    sheet.cell(14, 2, "Web outage")
    sheet.cell(14, 3, "Customer site becomes unavailable")
    populated_output = BytesIO()
    populated.save(populated_output)
    rows = import_register_workbook(populated_output.getvalue(), "scenario_risks")
    assert len(rows) == 1
    assert rows[0]["risk_ref"] == 1
    assert rows[0]["risk_description"] == "Customer site becomes unavailable"


def test_risk_register_normalizes_control_reference_lists_and_share_alias() -> None:
    """Keep pasted workbook values compatible with the persistent vocabulary."""

    rows = normalize_register_rows(
        "asset_risks",
        [
            {
                "risk_ref": 1,
                "asset": "Web server",
                "threat": "Service outage",
                "control_references": "A.8.20\nA.8.21",
                "treatment_option": "Share",
            }
        ],
    )
    assert rows[0]["control_references"] == ["A.8.20", "A.8.21"]
    assert rows[0]["treatment_option"] == "Transfer / Share"

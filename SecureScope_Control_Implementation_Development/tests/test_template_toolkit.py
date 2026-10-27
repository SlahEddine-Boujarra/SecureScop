"""Security and fidelity contracts for populated source toolkit workbooks."""

from __future__ import annotations

from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import pytest
from lxml import etree
from openpyxl import load_workbook

from securescope.reporting.template_toolkit import (
    TEMPLATE_SPECS,
    TemplateToolkitError,
    formula_cells,
    formula_error_count,
    import_soa_template,
    populate_toolkit_template,
    validate_recalculated_template,
)

PROJECT_ROOT = Path(__file__).resolve().parents[1]
TEMPLATE_ROOT = PROJECT_ROOT / "config" / "toolkit_templates"


def _sample_asset_row() -> dict[str, object]:
    """Return one valid risk row with formula-like free text."""
    return {
        "risk_ref": 1,
        "asset_group": "Information",
        "asset": "Customer database",
        "threat": "=scanner_formula()",
        "vulnerability": "Internet exposure",
        "risk_type": "C",
        "risk_owner": "CISO",
        "existing_controls": "Hardened server",
        "pre_likelihood": 3,
        "pre_likelihood_rationale": "Daily attempts",
        "pre_impact": 4,
        "pre_impact_rationale": "Material disruption",
        "treatment_option": "Modify",
        "treatment_action": "Patch and monitor",
        "control_references": ["A.8.7", "A.8.8"],
        "treatment_cost": 5000,
        "treatment_action_owner": "CISO",
        "treatment_action_timescale": "1 month",
        "treatment_action_progress": "Not started",
        "treatment_action_status": "Not Started",
        "post_likelihood": 2,
        "post_likelihood_rationale": "Lower after patching",
        "post_impact": 4,
        "post_impact_rationale": "Impact unchanged",
        "comments": "@user supplied note",
    }


def test_template_population_preserves_formula_and_nonworksheet_parts() -> None:
    """Only the target worksheet and recalculation metadata may change."""
    source = (TEMPLATE_ROOT / TEMPLATE_SPECS["asset_risks"].filename).read_bytes()
    populated = populate_toolkit_template(source, "asset_risks", [_sample_asset_row()])
    with ZipFile(BytesIO(source)) as original, ZipFile(BytesIO(populated)) as result:
        assert set(original.namelist()) == set(result.namelist())
        changed = {"xl/workbook.xml", "xl/worksheets/sheet3.xml"}
        for name in original.namelist():
            if name not in changed:
                assert original.read(name) == result.read(name), name
    assert formula_cells(source) == formula_cells(populated)
    assert formula_error_count(populated) == 0

    workbook = load_workbook(BytesIO(populated), read_only=True, data_only=False)
    row = workbook["Risk Assessment"]
    assert row["E14"].value == "'=scanner_formula()"
    assert row["N14"].value.startswith("=IFERROR")
    workbook.close()


def test_soa_template_reports_sample_values_that_need_human_justification() -> None:
    """The sample SoA is a preview source, not an approved client assessment."""
    source = (TEMPLATE_ROOT / TEMPLATE_SPECS["soa"].filename).read_bytes()
    rows, errors = import_soa_template(source)
    assert len(rows) == 93
    assert errors
    assert any("inclusion justification" in error for error in errors)
    assert all(row["approved"] is False for row in rows)


def test_soa_population_keeps_control_catalog_when_rows_are_partial() -> None:
    """A partial client export clears decisions but retains source control labels."""
    source = (TEMPLATE_ROOT / TEMPLATE_SPECS["soa"].filename).read_bytes()
    rows = [
        {
            "control_id": "A.5.1",
            "control_name": "Policies for information security",
            "applicability": "Applicable",
            "inclusion_justification": "Required by policy scope",
            "implementation_status": "Implemented",
            "effectiveness_review": "Reviewed",
        }
    ]
    populated = populate_toolkit_template(source, "soa", rows)
    workbook = load_workbook(BytesIO(populated), read_only=True, data_only=False)
    sheet = workbook["ISO27001 Annex A Applicability"]
    assert sheet["C15"].value == "A.5.1 Policies for information security"
    assert sheet["D15"].value == "Yes"
    assert sheet["E15"].value == "Yes"
    assert sheet["D16"].value is None
    assert sheet["E16"].value is None
    workbook.close()


def test_all_register_templates_accept_validated_rows_without_losing_formulas() -> None:
    """Scenario and opportunity templates use the same safe patch path."""
    rows_by_kind = {
        "scenario_risks": [
            {
                "risk_ref": 1,
                "risk_summary": "External access to customer service",
                "risk_description": "An internet-facing service is compromised.",
                "risk_type": "C",
                "risk_owner": "CISO",
                "pre_likelihood": 3,
                "pre_impact": 4,
                "treatment_option": "Modify",
                "treatment_action": "Patch and verify",
            }
        ],
        "opportunities": [
            {
                "ref": 1,
                "summary": "Improve secure delivery",
                "description": "A repeatable control review shortens release time.",
                "type": "Process",
                "owner": "CISO",
                "likelihood": 3,
                "impact": 4,
                "action": "Pilot the review checklist",
            }
        ],
    }
    for kind, rows in rows_by_kind.items():
        source = (TEMPLATE_ROOT / TEMPLATE_SPECS[kind].filename).read_bytes()
        populated = populate_toolkit_template(source, kind, rows)
        assert formula_cells(source) == formula_cells(populated)
        assert formula_error_count(populated) == 0


@pytest.mark.parametrize("kind", list(TEMPLATE_SPECS))
def test_population_preserves_excel_extension_namespace_bindings(kind: str) -> None:
    """Extension prefix names in Excel metadata must remain declared on save."""
    source = (TEMPLATE_ROOT / TEMPLATE_SPECS[kind].filename).read_bytes()
    populated = populate_toolkit_template(source, kind, [])
    ignorable = "{http://schemas.openxmlformats.org/markup-compatibility/2006}Ignorable"
    with ZipFile(BytesIO(populated)) as archive:
        for name in archive.namelist():
            if name.endswith(".xml"):
                root = etree.fromstring(archive.read(name))
                for node in root.iter():
                    assert all(
                        prefix in node.nsmap for prefix in node.get(ignorable, "").split()
                    ), name


def test_converter_copy_cannot_be_reported_as_recalculated() -> None:
    """A successful converter exit with stale source caches is insufficient."""
    source = (TEMPLATE_ROOT / TEMPLATE_SPECS["asset_risks"].filename).read_bytes()
    populated = populate_toolkit_template(source, "asset_risks", [_sample_asset_row()])
    with pytest.raises(TemplateToolkitError, match="stale|unavailable"):
        validate_recalculated_template(populated, populated)


def test_recalculation_rejects_loss_of_chart_objects() -> None:
    """A converter cannot silently drop a chart while retaining valid formulas."""
    source = (TEMPLATE_ROOT / TEMPLATE_SPECS["soa"].filename).read_bytes()
    output = BytesIO()
    with ZipFile(BytesIO(source)) as original, ZipFile(output, "w") as changed:
        for name in original.namelist():
            if name != "xl/charts/chart1.xml":
                changed.writestr(name, original.read(name))
    with pytest.raises(TemplateToolkitError, match="chart or pivot objects"):
        validate_recalculated_template(source, output.getvalue())


def test_import_rejects_xml_document_types() -> None:
    """Uploaded worksheet XML cannot resolve local or network entities."""
    source = (TEMPLATE_ROOT / TEMPLATE_SPECS["soa"].filename).read_bytes()
    output = BytesIO()
    with ZipFile(BytesIO(source)) as original, ZipFile(output, "w") as changed:
        for name in original.namelist():
            content = original.read(name)
            if name == "xl/workbook.xml":
                content = (
                    b'<!DOCTYPE workbook [<!ENTITY payload "unexpected">]>'
                    + content[content.index(b"?>") + 2 :]
                )
            changed.writestr(name, content)
    with pytest.raises(TemplateToolkitError, match="document type"):
        import_soa_template(output.getvalue())
